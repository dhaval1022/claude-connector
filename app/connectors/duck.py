"""CSV / Excel data sources, queried through a per-source read-only DuckDB database.

Accuracy rule: never let the sniffer guess. Every column is first loaded as text, then typed
only when ALL values fit a strict pattern. Identifier-like values (leading zeros such as
'00123', >18 digit numbers) stay text so they are never silently altered. Non-ISO dates stay
text with a note. Queries run with external file access disabled."""
import asyncio
import csv
import re
import threading
import time
from datetime import date, datetime
from pathlib import Path

import duckdb

from app.connectors.sql import _jsonable, _type_name

MAX_COLS = 500
INT_RE = r"^-?(0|[1-9][0-9]{0,17})$"
DEC_RE = r"^-?(0|[1-9][0-9]{0,17})\.[0-9]{1,10}$"
DATE_RE = r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$"
TS_RE = r"^[0-9]{4}-[0-9]{2}-[0-9]{2}[ T][0-9]{2}:[0-9]{2}(:[0-9]{2}(\.[0-9]{1,6})?)?$"
NONISO_DATE_RE = r"^[0-9]{1,2}[/.-][0-9]{1,2}[/.-][0-9]{2,4}$"


class FileLoadError(ValueError):
    """Safe-to-display problem with an uploaded file."""


def safe_ident(name: str, fallback: str) -> str:
    s = re.sub(r"[^a-z0-9_]+", "_", str(name).strip().lower()).strip("_")
    if not s:
        s = fallback
    return ("_" + s) if s[0].isdigit() else s[:63]


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, n in enumerate(names):
        base = safe_ident(n, f"col_{i + 1}")
        k = seen.get(base, 0)
        seen[base] = k + 1
        out.append(base if k == 0 else f"{base}_{k + 1}")
    return out


def _q(ident: str) -> str:
    return '"' + ident.replace('"', '""') + '"'


def _excel_to_csvs(path: Path, workdir: Path) -> list[tuple[str, Path, list[str]]]:
    """[(sheet name, csv path, notes)] - values written losslessly as text."""
    import openpyxl
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            notes: list[str] = []
            target = workdir / f"{safe_ident(path.stem, 'book')}__{safe_ident(ws.title, 'sheet')}.csv"
            rows = 0
            with open(target, "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                started = False
                for row in ws.iter_rows(values_only=True):
                    if not started and not any(c is not None and str(c).strip() for c in row):
                        continue  # skip blank rows above the header
                    started = True
                    cells = []
                    for c in row:
                        if c is None:
                            cells.append("")
                        elif isinstance(c, datetime):
                            cells.append(c.isoformat(sep=" ") if (c.hour or c.minute or c.second)
                                         else c.date().isoformat())
                        elif isinstance(c, date):
                            cells.append(c.isoformat())
                        elif isinstance(c, float) and c.is_integer() and abs(c) < 1e15:
                            cells.append(str(int(c)))
                        else:
                            cells.append(str(c))
                    w.writerow(cells)
                    rows += 1
            if rows == 0:
                target.unlink(missing_ok=True)
                continue
            if rows > 1 and ws.max_row and rows == 1:
                notes.append("Sheet has only a header row.")
            out.append((ws.title, target, notes))
    finally:
        wb.close()
    return out


def _load_csv_as_text(con, csv_path: Path, staging: str) -> list[str]:
    last = None
    for enc in ("utf-8", "latin-1"):
        try:
            con.execute(
                f"CREATE OR REPLACE TABLE {_q(staging)} AS SELECT * FROM read_csv(?, "
                "header=true, all_varchar=true, sample_size=-1, strict_mode=false, "
                "null_padding=true, parallel=false, encoding=?)", [str(csv_path), enc])
            return [r[0] for r in con.execute(f"DESCRIBE {_q(staging)}").fetchall()]
        except duckdb.Error as e:
            last = e
    raise FileLoadError(f"Could not read '{csv_path.name}' as CSV: {str(last)[:200]}")


def _infer_and_create(con, staging: str, final: str, raw_cols: list[str]) -> tuple[list, list]:
    """Create `final` from text staging table with strictly inferred types."""
    names = _dedupe(raw_cols)
    notes: list[str] = []
    selects = []
    for raw, name in zip(raw_cols, names):
        c = f"NULLIF(TRIM({_q(raw)}), '')"
        stats = con.execute(f"""
            SELECT count({c}),
                   count_if(regexp_matches({c}, '{INT_RE}')),
                   count_if(regexp_matches({c}, '{INT_RE}') OR regexp_matches({c}, '{DEC_RE}')),
                   max(CASE WHEN regexp_matches({c}, '{DEC_RE}')
                            THEN length(split_part({c}, '.', 2)) END),
                   count_if(regexp_matches({c}, '{DATE_RE}')
                            AND try_cast({c} AS DATE) IS NOT NULL),
                   count_if(regexp_matches({c}, '{TS_RE}')
                            AND try_cast({c} AS TIMESTAMP) IS NOT NULL),
                   count_if(regexp_matches({c}, '{NONISO_DATE_RE}')),
                   count_if(regexp_matches({c}, '^-?0[0-9]+$'))
            FROM {_q(staging)}""").fetchone()
        n, n_int, n_num, scale, n_date, n_ts, n_noniso, n_lead0 = stats
        if n == 0:
            selects.append(f"{c} AS {_q(name)}")
        elif n_int == n:
            selects.append(f"CAST({c} AS BIGINT) AS {_q(name)}")
        elif n_num == n:
            selects.append(f"CAST({c} AS DECIMAL(28,{max(int(scale or 0), 1)})) AS {_q(name)}")
        elif n_date == n:
            selects.append(f"CAST({c} AS DATE) AS {_q(name)}")
        elif n_ts == n or (n_date + n_ts == n):
            selects.append(f"CAST({c} AS TIMESTAMP) AS {_q(name)}")
        else:
            selects.append(f"{c} AS {_q(name)}")
            if n_lead0 and n_num + n_lead0 >= n * 0.9:
                notes.append(f"Column '{name}' has values with leading zeros; kept as text so "
                             "codes are not altered.")
            elif n_noniso == n:
                notes.append(f"Column '{name}' looks like dates in a non-ISO format "
                             "(e.g. dd/mm/yyyy); kept as text because day/month order is "
                             "ambiguous. Parse with strptime in SQL.")
    con.execute(f"CREATE OR REPLACE TABLE {_q(final)} AS SELECT {', '.join(selects)} "
                f"FROM {_q(staging)}")
    return names, notes


def build_database(db_path: Path, files: list[Path]) -> dict:
    """Load uploaded CSV/XLSX files into a fresh DuckDB file; returns the schema cache."""
    db_path.unlink(missing_ok=True)
    work = db_path.parent / "_work"
    work.mkdir(exist_ok=True)
    con = duckdb.connect(str(db_path))
    tables: dict = {}
    try:
        for f in files:
            ext = f.suffix.lower()
            if ext == ".csv" or ext == ".txt":
                units = [(None, f, [])]
            elif ext in (".xlsx", ".xlsm"):
                try:
                    units = [(s, p, n) for s, p, n in _excel_to_csvs(f, work)]
                except Exception as e:
                    raise FileLoadError(f"Could not read Excel file '{f.name}': {str(e)[:200]}")
                if not units:
                    raise FileLoadError(f"'{f.name}' contains no data.")
            else:
                raise FileLoadError(f"Unsupported file type '{ext}'. Use .csv or .xlsx.")
            for sheet, csv_path, notes in units:
                base = safe_ident(f.stem if sheet is None else f"{f.stem}_{sheet}", "table")
                final, k = base, 2
                while final in tables:
                    final, k = f"{base}_{k}", k + 1
                raw_cols = _load_csv_as_text(con, csv_path, "_staging")
                if len(raw_cols) > MAX_COLS:
                    raise FileLoadError(f"'{f.name}' has more than {MAX_COLS} columns.")
                names, col_notes = _infer_and_create(con, "_staging", final, raw_cols)
                rows = con.execute(f"SELECT count(*) FROM {_q(final)}").fetchone()[0]
                types = con.execute(f"DESCRIBE {_q(final)}").fetchall()
                tables[final] = {
                    "kind": "table", "primary_key": [], "foreign_keys": [], "rows": rows,
                    "source_file": f.name + (f" [sheet: {sheet}]" if sheet else ""),
                    "notes": notes + col_notes,
                    "columns": [{"name": t[0], "type": t[1], "nullable": True} for t in types]}
        con.execute("DROP TABLE IF EXISTS _staging")
        con.execute("CHECKPOINT")
    finally:
        con.close()
        for p in work.glob("*"):
            p.unlink(missing_ok=True)
        work.rmdir()
    if not tables:
        raise FileLoadError("No tables could be loaded.")
    return {"default_schema": "main", "schemas": {"main": tables}}


class DuckConnector:
    dialect = "duckdb"
    engine_name = "files"

    def __init__(self, source_id: str, cfg: dict):
        self.source_id, self.cfg = source_id, cfg
        self.db_path = cfg["db_path"]

    def _open(self):
        # external access off => read_csv / COPY / ATTACH / file reads fail at engine level
        return duckdb.connect(self.db_path, read_only=True, config={
            "enable_external_access": False, "threads": 2, "memory_limit": "1GB"})

    def _execute(self, sql: str, max_rows: int, timeout_s: int) -> dict:
        start = time.monotonic()
        con = self._open()
        timer = threading.Timer(timeout_s, con.interrupt)
        timer.start()
        try:
            cur = con.execute(sql)
            columns = [d[0] for d in cur.description]
            raw = cur.fetchmany(max_rows + 1)
        except duckdb.InterruptException:
            raise TimeoutError(f"Query exceeded the {timeout_s}s time limit.")
        finally:
            timer.cancel()
            con.close()
        truncated = len(raw) > max_rows
        raw = raw[:max_rows]
        cols_vals = list(zip(*raw)) if raw else [[] for _ in columns]
        return {"columns": [{"name": c, "type": _type_name(list(v))}
                            for c, v in zip(columns, cols_vals)],
                "rows": [[_jsonable(v) for v in r] for r in raw],
                "row_count": len(raw), "truncated": truncated,
                "elapsed_ms": int((time.monotonic() - start) * 1000)}

    def stream(self, sql: str, timeout_s: int, batch: int = 5000):
        con = self._open()
        timer = threading.Timer(timeout_s, con.interrupt)
        timer.start()
        try:
            cur = con.execute(sql)
            yield {"columns": [d[0] for d in cur.description], "types": None}
            while True:
                rows = cur.fetchmany(batch)
                if not rows:
                    break
                yield {"rows": rows}
        except duckdb.InterruptException:
            raise TimeoutError(f"Query exceeded the {timeout_s}s time limit.")
        finally:
            timer.cancel()
            con.close()

    async def execute(self, sql: str, max_rows: int, timeout_s: int) -> dict:
        return await asyncio.to_thread(self._execute, sql, max_rows, timeout_s)

    async def introspect(self) -> dict:
        return self.cfg["schema"]  # fixed at load time; files are immutable

    async def test(self) -> None:
        await self.execute("SELECT 1", 1, 10)

    @staticmethod
    def dispose(source_id: str) -> None:
        pass
